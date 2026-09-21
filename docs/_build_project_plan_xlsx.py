"""Build docs/SRSE_PROJECT_PLAN.xlsx — run: python3.12 docs/_build_project_plan_xlsx.py (venv + openpyxl)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent / "SRSE_PROJECT_PLAN.xlsx"

HEADER_FONT = Font(name="Arial", bold=True, size=11)
BODY_FONT = Font(name="Arial", size=11)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
WRAP = Alignment(wrap_text=True, vertical="top")

# sno, activity, pct (0–1), dependency, target date, remarks
PLAN_ROWS = [
    (
        1,
        "Business Requirement Document (BRD)",
        0.95,
        "SMART programme inputs",
        "",
        "docs/SRSE_BRD.docx v1.1 — stakeholder narrative, workflow prose, config backup §10; aligned with TDD v1.2. Pending sign-off.",
    ),
    (
        2,
        "Design Document",
        0.98,
        "1",
        "",
        "docs/SRSE_Technical_Design_Document.docx v1.2 — stakeholder overview, service catalog, JSON config, Analysis engine; PPT Part B.",
    ),
    (
        3,
        "Platform scaffold",
        1.00,
        "2",
        "",
        "Docker Compose, DB2 (operational) + Presto (analytical), and DATA_MODE synthetic/live switch proving both planes connect.",
    ),
    (
        4,
        "Rule Engine",
        0.95,
        "3",
        "",
        "DMN-shaped AST, SQL compiler with bound parameters only, unit-tested; maps officer fields to physical columns via metadata.",
    ),
    (
        5,
        "Execution service",
        0.90,
        "4",
        "",
        "Push-down count/breakdown/cohort against Presto; query timeout, aggregate-only results, fixed breakdown dimensions, cohort hard cap.",
    ),
    (
        6,
        "Metadata & mapping",
        0.90,
        "3",
        "",
        "Field catalogue and per-mode column mappings in DB2; admin editor; Caffeine-cached resolver for rule and preview paths.",
    ),
    (
        7,
        "Lakehouse registry & Analysis metadata",
        0.90,
        "6",
        "",
        "Register Silver/Gold tables; live column introspection; compare-as and fuzzy flags for cross-type and name matching.",
    ),
    (
        8,
        "Officer Rules UI",
        0.85,
        "4, 6",
        "",
        "Rule builder, threshold sliders/selects, simulation results, district/gender/age breakdown tables and charts (e.g. Ekal Naari).",
    ),
    (
        9,
        "Scenario store & comparison",
        0.75,
        "4, 8",
        "",
        "Save ruleset + result snapshots; load and pairwise compare breakdown deltas when thresholds change.",
    ),
    (
        10,
        "Cohort drill-down API",
        0.80,
        "5",
        "",
        "Row-level sample endpoint with server-enforced cap; reports truncation so samples are not mistaken for full cohorts.",
    ),
    (
        11,
        "Cohort drill-down UI",
        0.40,
        "10",
        "",
        "Officer-facing drill-down from simulation results; API exists; UI wiring and UX still limited.",
    ),
    (
        12,
        "Analysis: single-table match",
        1.00,
        "7",
        "",
        "Two-table join on registered columns; exact/fuzzy match; NDJSON stream and full CSV download from Presto.",
    ),
    (
        13,
        "Analysis: display-only columns",
        1.00,
        "12",
        "",
        "Extra columns in result grid only—not used in join keys—so officers see context without changing match logic.",
    ),
    (
        14,
        "Analysis: multi-target hub match",
        1.00,
        "13",
        "",
        "One hub vs N targets (separate joins); progress events; time budget; partial failure per target; CSV all-or-nothing.",
    ),
    (
        15,
        "Analysis: column groups (COMBINE / ANY_OF)",
        0.95,
        "13",
        "",
        "Multi-column criteria, name folding, UNNEST-based ANY_OF; SQL checked against Presto 0.297 parser in tests.",
    ),
    (
        16,
        "Analysis: manual E2E on Presto",
        0.80,
        "15",
        "",
        "Container-based cross-catalog runs when match SQL changes; not fully automated in CI.",
    ),
    (
        17,
        "Admin (lakehouse, mapping, backup)",
        0.95,
        "6, 7",
        "",
        "Table registration, field mapping editor, scheme settings, JSON export/import for operational config backup.",
    ),
    (
        18,
        "Configuration files",
        0.88,
        "3",
        "",
        "Runtime config: application.yml, Spring profiles, guardrail/analysis properties, JDBC URLs, DATA_MODE and auth mode; docker-compose.yml and client-dev overrides; env-var catalogue for deploy.",
    ),
    (
        19,
        "Configuration guide (documentation)",
        0.85,
        "18",
        "",
        "Human-readable ops doc: what each setting does, local vs client-dev values, troubleshooting (CONFIGURATION_GUIDE).",
    ),
    (
        20,
        "Code quality (SonarQube)",
        1.00,
        "—",
        "",
        "Java and TypeScript issues addressed for maintainability and CI quality gate.",
    ),
    (
        21,
        "Frontend lint cleanup",
        0.30,
        "—",
        "",
        "Remaining ESLint errors in legacy components outside Analysis; build passes but lint debt remains.",
    ),
    (
        22,
        "Test cases",
        0.85,
        "4–16",
        "",
        "docs/SRSE_Test_Cases.xlsx — 40 SIT/UAT cases with BRD/TDD traceability; automated backend suite (~300+ unit/integration tests) separate.",
    ),
    (
        23,
        "SIT",
        0.25,
        "22, 25–27",
        "",
        "End-to-end integration on synthetic stack and client-dev Presto/DB2: rules, Analysis, admin, auth seams.",
    ),
    (
        24,
        "Client Dev profile",
        0.70,
        "3, 18",
        "",
        "Compose profile and settings to point at on-prem lakehouse and DB2; verified in slices; full stack boot sign-off pending.",
    ),
    (
        25,
        "Presto JDBC confirmation",
        0.00,
        "Lovadeep / IBM",
        "",
        "Confirm PrestoDB (not Trino) driver coordinates for watsonx.data 2.3.1 against live endpoint.",
    ),
    (
        26,
        "Golden Layer naming & mappings",
        0.30,
        "Lovadeep / DBA",
        "",
        "Replace placeholders with real catalog.schema.table.column bindings for LIVE mode.",
    ),
    (
        27,
        "SRSE DB2 placement",
        0.50,
        "Lovadeep",
        "",
        "Final decision on operational schema/instance for SRSE-owned JPA entities.",
    ),
    (
        28,
        "Derived fields (REL-01, etc.)",
        0.40,
        "Lovadeep / data eng.",
        "",
        "Confirm Tier-3 fields pre-materialised upstream (e.g. income 3-yr avg) before officers map them.",
    ),
    (
        29,
        "Live lakehouse seed & validation",
        0.40,
        "26, 28",
        "",
        "Register production tables, complete mappings, run representative scheme simulations on live data.",
    ),
    (
        30,
        "Auth (RajSewadwar SSO)",
        0.20,
        "Arvind",
        "",
        "Mock JWT and STATE_OFFICER gate locally; production SSO payload parsing still stubbed, fails closed.",
    ),
    (
        31,
        "Security review",
        0.00,
        "24, 30",
        "",
        "Client-environment review of auth, injection controls, and deployment surface.",
    ),
    (
        32,
        "UAT",
        0.00,
        "23, 29, 30",
        "",
        "Departmental officers validate rules and Analysis workflows against agreed test cases.",
    ),
    (
        33,
        "Deployment",
        0.60,
        "23, 24",
        "",
        "Two-container packaging, install steps, smoke checks, runbooks; production monitoring hooks TBD.",
    ),
    (
        34,
        "Services inventory",
        0.75,
        "2",
        "",
        "Master list of REST capabilities and UI modules—developed vs planned—see Reference for row 34 tab.",
    ),
]

SERVICES = [
    ("Health & connectivity", "Developed", "HealthController, connection info/update"),
    ("Decision / rules simulation", "Developed", "/api/decision/**"),
    ("Presto execution", "Developed", "Count, breakdown, cohort (via decision layer)"),
    ("Field catalogue & mapping", "Developed", "FieldCatalogController, FieldColumnMappingController"),
    ("Scheme configuration", "Developed", "SchemeController"),
    ("Scenario store & compare", "Developed", "Scenario APIs under decision"),
    ("Lakehouse admin & browse", "Developed", "LakehouseAdminController, LakehouseCatalogController"),
    ("Analysis column metadata", "Developed", "AnalysisColumnMetadataController"),
    ("Record match / Analysis", "Developed", "RecordMatchController (single + multi, CSV)"),
    ("Admin config backup", "Developed", "AdminConfigController"),
    ("Officer UI", "Developed (partial cohort UI)", "Next.js frontend container"),
    ("RajSewadwar SSO", "To be developed", "Production auth"),
    ("Live Golden Layer catalogue", "To be developed", "Depends on DBA inputs"),
    ("Formal SIT sign-off pack", "To be developed", "Row 23 on Project Plan"),
    ("Production handover", "To be developed", "Row 33 on Project Plan"),
]

REF_SHEET = "Reference for row 34"


def style_header_row(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Project Plan"

    headers = [
        "S. No.",
        "Activity",
        "% Completed",
        "Dependency",
        "Target Completion Date",
        "Remarks",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h)
    style_header_row(ws, 1, len(headers))

    first_data = 2
    last_data = first_data + len(PLAN_ROWS) - 1
    for i, row in enumerate(PLAN_ROWS, start=first_data):
        sno, activity, pct, dep, target, remarks = row
        ws.cell(row=i, column=1, value=sno).font = BODY_FONT
        ac = ws.cell(row=i, column=2, value=activity)
        ac.font = BODY_FONT
        ac.alignment = WRAP
        pc = ws.cell(row=i, column=3, value=pct)
        pc.font = BODY_FONT
        pc.number_format = "0%"
        ws.cell(row=i, column=4, value=dep).font = BODY_FONT
        ws.cell(row=i, column=4).alignment = WRAP
        ws.cell(row=i, column=5, value=target).font = BODY_FONT
        rc = ws.cell(row=i, column=6, value=remarks)
        rc.font = BODY_FONT
        rc.alignment = WRAP

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 52
    ws.freeze_panes = "A2"

    ws_ref = wb.create_sheet(REF_SHEET)
    ref_headers = ["Service / capability", "Status", "Primary API / surface"]
    for col, h in enumerate(ref_headers, start=1):
        ws_ref.cell(row=1, column=col, value=h)
    style_header_row(ws_ref, 1, len(ref_headers))
    for i, (svc, status, api) in enumerate(SERVICES, start=2):
        ws_ref.cell(row=i, column=1, value=svc).font = BODY_FONT
        ws_ref.cell(row=i, column=1).alignment = WRAP
        ws_ref.cell(row=i, column=2, value=status).font = BODY_FONT
        ws_ref.cell(row=i, column=3, value=api).font = BODY_FONT
        ws_ref.cell(row=i, column=3).alignment = WRAP
    ws_ref.column_dimensions["A"].width = 42
    ws_ref.column_dimensions["B"].width = 28
    ws_ref.column_dimensions["C"].width = 52
    ws_ref.freeze_panes = "A2"

    ws_sum = wb.create_sheet("Overall completion")
    ws_sum.cell(row=1, column=1, value="Metric").font = HEADER_FONT
    ws_sum.cell(row=1, column=2, value="Value").font = HEADER_FONT
    style_header_row(ws_sum, 1, 2)

    pct_col = f"'Project Plan'!C{first_data}:C{last_data}"
    ws_sum.cell(row=2, column=1, value="Overall % completed (simple average of all activities)").font = BODY_FONT
    ws_sum.cell(row=2, column=1).alignment = WRAP
    overall = ws_sum.cell(row=2, column=2, value=f"=AVERAGE({pct_col})")
    overall.font = Font(name="Arial", bold=True, size=12)
    overall.number_format = "0.0%"

    ws_sum.cell(row=4, column=1, value="Activities counted").font = BODY_FONT
    ws_sum.cell(row=4, column=2, value=f"=COUNT({pct_col})")
    ws_sum.cell(row=4, column=2).font = BODY_FONT

    ws_sum.cell(row=5, column=1, value="Note").font = BODY_FONT
    note = ws_sum.cell(
        row=5,
        column=2,
        value="Unweighted mean of % Completed on Project Plan tab. Update activity rows to refresh.",
    )
    note.font = Font(name="Arial", size=10, italic=True)
    note.alignment = WRAP
    ws_sum.column_dimensions["A"].width = 48
    ws_sum.column_dimensions["B"].width = 36

    wb.save(OUT)
    print(f"Wrote {OUT} ({len(PLAN_ROWS)} activities)")


if __name__ == "__main__":
    main()
