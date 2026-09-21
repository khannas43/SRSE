"""Build docs/SRSE_BRD.docx from brd_assets workflow PNGs."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "brd_assets"
OUT = ROOT / "SRSE_BRD.docx"


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    if bold:
        run.bold = True


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(11)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(10)
    doc.add_paragraph()


def add_figure(doc: Document, png: str, caption: str) -> None:
    path = ASSETS / png
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.2))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:
            r.italic = True
            r.font.size = Pt(10)
            r.font.name = "Arial"
    else:
        para(doc, f"[Missing diagram: {png}]")
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    # Title block
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Business Requirements Document (BRD)")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)
    r.font.name = "Arial"

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(
        "Scheme Rule Simulation Engine (SRSE)\n"
        "Standalone eligibility simulation and data-reconciliation product"
    )
    sr.font.size = Pt(14)
    sr.font.name = "Arial"

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        f"Document version: 1.1\nDate: {date.today():%d %B %Y}\nStatus: For stakeholder review\n"
        "Editable workflow & architecture slides: docs/SRSE_Functional_Workflows.pptx (Part A workflows)\n"
        "Technical detail: docs/SRSE_Technical_Design_Document.docx v1.2\n"
        "(First deployed under Rajasthan SMART; readable without prior SMART programme knowledge.)"
    )
    mr.font.size = Pt(10)
    mr.font.name = "Arial"
    doc.add_page_break()

    heading(doc, "Document control", 1)
    add_table(
        doc,
        ["Version", "Date", "Author", "Description"],
        [
            ["0.1", "2026-Q1", "Delivery team", "Initial scope from TDD"],
            ["1.0", "2026-09-20", "SRSE team", "Complete BRD with functional workflows"],
            [
                "1.1",
                date.today().isoformat(),
                "SRSE team",
                "Stakeholder narrative, workflow prose, config backup & delivery activities; aligned with TDD v1.2",
            ],
        ],
    )

    heading(doc, "Distribution", 2)
    add_table(
        doc,
        ["Role", "Organisation", "Purpose"],
        [
            ["Programme owner", "DoIT&C / RISL", "Approval"],
            ["Department officers", "Line departments", "UAT & adoption"],
            ["Lakehouse admin", "Data platform team", "Registration & metadata"],
            ["Implementation", "Deloitte / SI", "Build & deploy"],
        ],
    )

    heading(doc, "1. Executive summary", 1)
    para(
        doc,
        "The Scheme Rule Simulation Engine (SRSE) is an independent software product that helps "
        "government departments answer two kinds of question quickly and repeatably: How many "
        "citizens would qualify for a scheme if we changed eligibility rules or thresholds? and "
        "How well do two beneficiary datasets align with each other? SRSE is delivered as its own "
        "web application and API. It does not require the SMART citizen portal to function and "
        "can be evaluated on its own merits by auditors, integrators, or other states.",
    )
    para(
        doc,
        "In the Rajasthan SMART programme, SRSE connects to the state lakehouse (Presto over "
        "Iceberg tables) for population-scale counts and to a small operational database for "
        "configuration, saved scenarios, and governance metadata. Departmental officers use a "
        "guided interface—no SQL required—to build rules from approved citizen attributes, preview "
        "results, save named scenarios for comparison, and run data-reconciliation matches "
        "between registered tables.",
    )
    para(
        doc,
        "Business value: faster policy what-if analysis, auditable scenario history, reduced "
        "dependence on ad-hoc scripts and spreadsheets, and a controlled path for reconciling "
        "source registers with golden beneficiary data. Privacy is preserved by returning "
        "aggregates in normal use; any row-level sample is strictly capped.",
    )
    bullet(doc, "Rules simulation — counts and standard breakdowns (district, gender, age band).")
    bullet(doc, "Analysis — exact/fuzzy matching, multi-column names, multi-target hub reconciliation.")
    bullet(doc, "Administration — register lakehouse tables, map fields, export/import JSON configuration.")
    bullet(doc, "Independent release — own containers, not tied to portal deployment cycles.")

    heading(doc, "2. Background and problem statement", 1)
    para(
        doc,
        "Welfare and subsidy schemes depend on eligibility conditions that combine age, income, "
        "location, social category, and many other attributes. When policy makers propose a "
        "change—such as raising an income ceiling—they need a defensible estimate of impact before "
        "issuing orders. Traditional approaches rely on IT teams to write one-off database queries, "
        "export CSV files, and manipulate them in Excel. Each iteration takes days or weeks, "
        "results are hard to reproduce, and knowledge leaves the organisation when staff move.",
    )
    para(
        doc,
        "Separately, departments accumulating data in multiple systems need to know whether the "
        "same person appears in two registers under slightly different spellings or ID formats. "
        "Manual spot checks do not scale to statewide populations.",
    )
    para(
        doc,
        "SRSE addresses both problems in one product. Rules simulation encodes eligibility as "
        "transparent, reusable logic against a curated field catalogue. Analysis provides "
        "governed matching between administrator-approved tables, including fuzzy name matching "
        "and large-result CSV export. The Rajasthan SMART lakehouse is the first analytical "
        "data source; the business requirements herein remain valid for any equivalent Presto/Iceberg "
        "deployment that meets the integration assumptions in Section 6.",
    )

    heading(doc, "3. Business objectives", 1)
    add_table(
        doc,
        ["ID", "Objective", "Success measure"],
        [
            ["BO-1", "Officers can simulate eligibility without IT ticket for each change", "Preview run in UI < agreed timeout"],
            ["BO-2", "Results are aggregate-only by default (privacy)", "No row-level data except capped cohort API"],
            ["BO-3", "Scenarios are saved and comparable", "Named snapshots with Δ breakdown"],
            ["BO-4", "Analysis supports Silver↔Gold reconciliation", "Match + CSV at crore scale via download"],
            ["BO-5", "Admin governs which lakehouse tables officers may use", "Registration allow-list enforced"],
        ],
    )

    heading(doc, "4. Scope", 1)
    heading(doc, "4.1 In scope", 2)
    bullet(doc, "Web UI: Rules builder, simulation results, Analysis tab, Admin (lakehouse, mappings, backup).")
    bullet(doc, "REST APIs: Decision preview/scenarios/cohort; Analysis match (single, multi-target, CSV).")
    bullet(doc, "Metadata: Field catalogue, field→column mapping per DATA_MODE, analysis column metadata.")
    bullet(doc, "Deployment: Two application containers (frontend + backend); synthetic and live data modes.")
    bullet(doc, "Authentication seam: STATE_OFFICER for decision paths; admin for lakehouse browse/register.")

    heading(doc, "4.2 Out of scope", 2)
    bullet(doc, "Beneficiary application processing or payment disbursement (portal / core schemes).")
    bullet(doc, "Authoring rules inside IBM CP4BA/ODM (future swap behind REST seam only).")
    bullet(doc, "On-the-fly JOINs or cross-table expressions in the rule compiler (Tier-3 fields must be pre-materialised upstream).")
    bullet(doc, "Unbounded export of full beneficiary populations.")

    heading(doc, "5. Stakeholders and user roles", 1)
    add_table(
        doc,
        ["Role", "Description", "Primary activities"],
        [
            ["State officer (STATE_OFFICER)", "Department user designing/simulating schemes", "Rules, scenarios, Analysis match"],
            ["Lakehouse administrator", "Data platform / SMART admin", "Register tables, column metadata, field mappings"],
            ["Programme owner", "DoIT&C / RISL", "Prioritisation, acceptance"],
            ["DBA / data engineering", "Golden Layer owners", "Physical columns, derived fields (REL-01)"],
        ],
    )

    heading(doc, "6. Assumptions and dependencies", 1)
    bullet(doc, "PrestoDB 0.297 (watsonx.data 2.3.1 lineage) is the analytical engine; PrestoDB JDBC driver confirmed for production.")
    bullet(doc, "Golden Layer exposes flat columns for all officer-visible fields (no runtime joins in SRSE).")
    bullet(doc, "RajSewadwar SSO provides identity and STATE_OFFICER role in production.")
    bullet(doc, "Operational DB2 instance/schema provisioned for SRSE entities.")
    bullet(doc, "Reference scheme criteria available (e.g. Ekal Naari pension rules from consolidated criteria workbook).")

    heading(doc, "7. Functional requirements", 1)

    heading(doc, "7.1 Rules and simulation", 2)
    add_table(
        doc,
        ["ID", "Requirement", "Priority"],
        [
            ["FR-R1", "Officer shall compose rules from allow-listed catalogue fields only", "Must"],
            ["FR-R2", "System shall compile rules to parameterised SQL and execute count in Presto", "Must"],
            ["FR-R3", "Preview shall return total count and optional breakdown (district, gender, age band)", "Must"],
            ["FR-R3a", "Rules UI default preview/cohort sample size configurable (SRSE_PREVIEW_SAMPLE_SIZE); clamped to SRSE_COHORT_CAP", "Must"],
            ["FR-R4", "Officer shall adjust thresholds (numeric ranges, IN lists, boolean toggles) without code change", "Must"],
            ["FR-R5", "Cohort sample shall be hard-capped (SRSE_COHORT_CAP); response states truncation", "Must"],
            ["FR-R8", "Scheme official template: officer save forks to new scenario; only admin sets template", "Should"],
            ["FR-R6", "Officer shall save a named scenario linked to one or more schemes with result snapshot", "Should"],
            ["FR-R7", "Officer shall compare two scenarios (count delta + breakdown deltas)", "Should"],
        ],
    )

    heading(doc, "7.2 Analysis (record match)", 2)
    para(
        doc,
        "Analysis supports data stewards and officers who must reconcile two datasets—not simulate "
        "a single ruleset. Users select registered source and target tables, define match keys "
        "(including combined name fields or match-if-any-of columns), and review overlapping and "
        "non-overlapping records. Fuzzy matching helps where spelling differs between systems. "
        "Multi-target mode allows one central register to be compared against several targets in "
        "one session (for example golden layer versus multiple departmental sources). "
        "Join-key suggestions help pick candidate columns (metadata-first; optional overlap probe "
        "samples the source and scans the target in full per pair).",
    )
    add_table(
        doc,
        ["ID", "Requirement", "Priority"],
        [
            ["FR-A1", "Officer shall match two registered tables on chosen column pairs or column groups", "Must"],
            ["FR-A2", "System shall support exact and fuzzy (Levenshtein) matching per group policy", "Must"],
            ["FR-A3", "System shall support COMBINE and ANY_OF column groups with documented SQL caps", "Must"],
            ["FR-A4", "Officer may add display-only columns not used in join keys", "Must"],
            ["FR-A5", "Multi-target mode: one hub table vs N targets (separate joins, combined grid)", "Must"],
            ["FR-A6", "Results stream as NDJSON; CSV download re-runs full query (not UI buffer)", "Must"],
            ["FR-A7", "UI shall defer grid/charts above 10,000 rows and offer CSV instead", "Should"],
            ["FR-A8", "Two-table match shall support INNER/LEFT/RIGHT/FULL join types and SQL preview", "Must"],
            ["FR-A9", "Multi-target join canvas serialises to same request as form (Phase A); INNER only", "Should"],
            ["FR-A10", "Join-key suggestions with optional overlap probe (sample source, full target scan)", "Should"],
        ],
    )

    heading(doc, "7.3 Administration and metadata", 2)
    add_table(
        doc,
        ["ID", "Requirement", "Priority"],
        [
            ["FR-M1", "Admin shall browse live lakehouse catalog (admin-only) and register tables for officers", "Must"],
            ["FR-M2", "Registered table columns are re-read live; hidden columns excluded from officer pickers", "Must"],
            ["FR-M3", "Admin shall set analysis column metadata (business name, fuzzy default, compare-as)", "Must"],
            ["FR-M4", "Admin shall maintain field catalogue and physical mappings for synthetic and live modes", "Must"],
            ["FR-M5", "Admin shall export/import JSON backup of operational configuration", "Should"],
        ],
    )

    heading(doc, "7.4 Security and access", 2)
    add_table(
        doc,
        ["ID", "Requirement", "Priority"],
        [
            ["FR-S1", "Decision APIs require authenticated STATE_OFFICER in production", "Must"],
            ["FR-S2", "Catalog/schema/table names validated against lakehouse allow-lists before SQL interpolation", "Must"],
            ["FR-S3", "Officer column picks validated against registration + live column existence", "Must"],
            ["FR-S4", "Mock JWT issuer available for local/dev only", "Should"],
            ["FR-S5", "SRSE_ADMIN vs STATE_OFFICER enforced per API path; mock mode not production ACL", "Must"],
        ],
    )

    heading(doc, "8. Non-functional requirements", 1)
    add_table(
        doc,
        ["ID", "Category", "Requirement"],
        [
            ["NFR-1", "Performance", "Count/breakdown queries respect configurable Presto timeout"],
            ["NFR-2", "Scalability", "Analysis match uncapped server-side; large results via CSV stream"],
            ["NFR-3", "Availability", "Health endpoint reports operational + analytical plane status"],
            ["NFR-4", "Maintainability", "Java 17 / Spring Boot 3; Next.js 16 frontend; containerised deploy"],
            ["NFR-5", "Auditability", "Scenarios persist ruleset JSON and result snapshots in DB2"],
            ["NFR-6", "Privacy", "Aggregate-only simulation; cohort cap on row-level drill-down"],
        ],
    )

    heading(doc, "9. Business rules and constraints", 1)
    bullet(doc, "Flat catalogue: officers never author joins; Tier-3 derived attributes must exist as Golden Layer columns.")
    bullet(doc, "Four-part lakehouse addressing: catalog.schema.table.column in LIVE bindings; layer is a registration/display tag — officer cascade filters by layer but payloads use catalog.schema.table only.")
    bullet(doc, "Bronze/Silver/Gold (and legacy untagged) tables may be registered; new registrations require a layer tag.")
    bullet(doc, "Cross-type comparisons use shared coercion rules (TRY_CAST on text for numeric identifiers by default).")
    bullet(doc, "Multi-target Analysis is N two-table INNER joins, never one N-way join; join canvas is a UI over the same engine.")
    bullet(doc, "DATA_MODE=synthetic|live switches resolver and mappings without code change; UI labels (Development/UAT/Production) and SRSE_ENV_LABEL are display-only.")
    bullet(doc, "Admin mapping editor SYNTHETIC/LIVE radios select binding set, not deployment environment.")

    heading(doc, "10. Configuration backup (business view)", 1)
    para(
        doc,
        "Administrators can export the full SRSE operational configuration as a single JSON file "
        "and re-import it after redeployment or disaster recovery. The bundle includes JDBC "
        "connection settings (treat exports as confidential—they may contain passwords), the "
        "officer field catalogue, mappings from business field names to lakehouse columns, "
        "registered tables, analysis column labels, and scheme definitions.",
    )
    para(
        doc,
        "Import merges by natural keys (field codes, table names, scheme codes)—it does not "
        "require the same internal database ids as the source environment. This supports promotion "
        "from test to production and backup before major mapping changes. Technical schema detail "
        "is documented in the Technical Design Document Section 7.",
    )

    heading(doc, "11. Functional workflows", 1)
    para(
        doc,
        "The workflows below describe how people and systems interact to deliver SRSE outcomes. "
        "Each subsection includes narrative for non-technical readers and a diagram. Workshop "
        "facilitators may edit the same flows in docs/SRSE_Functional_Workflows.pptx (Part A).",
    )

    heading(doc, "11.1 System context", 2)
    para(
        doc,
        "End users reach SRSE through a standard web browser. They do not connect directly to "
        "the lakehouse or operational database. The SRSE application forwards authorised requests "
        "to the API layer, which in turn reads configuration from the operational store and "
        "runs analytical work on the lakehouse query engine.",
    )
    para(
        doc,
        "Identity is established through the organisation's single sign-on. Only users with the "
        "appropriate role may run rule simulations or open administrative functions. This boundary "
        "ensures that citizen-scale data remains behind authenticated, role-checked services.",
    )
    add_figure(doc, "wf01_system_context.png", "Figure 11-1: SRSE in the enterprise landscape")
    para(
        doc,
        "Figure 11-1 should be read as a logical view: SRSE may be hosted in the same data centre "
        "as the lakehouse but remains a separately deployable product.",
    )

    heading(doc, "11.2 Admin — lakehouse and metadata setup", 2)
    para(
        doc,
        "Before officers can simulate rules or run Analysis, a data administrator registers which "
        "lakehouse tables SRSE may use and maps business-friendly field names to physical columns. "
        "Registration is deliberate—officers cannot query arbitrary tables on the cluster. Column "
        "lists are refreshed from the live lakehouse so new upstream columns appear without "
        "re-registration, while removed columns disappear from pickers.",
    )
    para(
        doc,
        "Administrators may also set display names, default fuzzy-matching behaviour, and how "
        "dissimilar data types should compare (for example numeric IDs stored as text). Optional "
        "JSON export captures this work for backup or migration.",
    )
    add_figure(doc, "wf02_admin_onboarding.png", "Figure 11-2: Administrator onboarding flow")
    para(
        doc,
        "This workflow is typically performed once per environment (or when new golden-layer "
        "tables are published), not by every officer.",
    )

    heading(doc, "11.3 Officer — rule simulation (preview)", 2)
    para(
        doc,
        "The core officer journey is iterative policy design. The user selects a welfare scheme "
        "context, builds eligibility from the flat field catalogue (logical AND/OR groups), and "
        "adjusts thresholds such as maximum income or eligible age range. Each preview request "
        "returns an updated statewide count without saving unless the user chooses to.",
    )
    para(
        doc,
        "Breakdowns by district, gender, and age band help departments understand geographic and "
        "demographic distribution of impact. Because preview does not persist data, officers can "
        "experiment freely before committing a scenario name for official comparison.",
    )
    add_figure(doc, "wf03_rule_simulation.png", "Figure 11-3: Rule build and preview flow")
    para(
        doc,
        "Downstream budgeting and cabinet notes can cite saved scenarios once stakeholders agree "
        "on a baseline and alternative thresholds.",
    )

    heading(doc, "11.4 Scenario save and compare", 2)
    para(
        doc,
        "When an officer saves a scenario, SRSE stores the ruleset and a snapshot of the "
        "simulation results at that moment. Naming scenarios supports workshop discussions "
        "(for example Baseline 2026 vs Income cap +5000).",
    )
    para(
        doc,
        "Compare mode takes two saved scenarios and reports the difference in total eligible "
        "count and in each breakdown dimension. This answers whether a proposed change materially "
        "shifts beneficiaries between districts or categories—not merely whether the headline number moved.",
    )
    add_figure(doc, "wf04_scenario_lifecycle.png", "Figure 11-4: Scenario lifecycle")

    heading(doc, "11.5 Analysis — record match", 2)
    para(
        doc,
        "Analysis workflows begin with selecting two registered tables (or one hub and multiple "
        "targets). The user defines match criteria, runs the job, and reviews results in the "
        "application or downloads a complete CSV when volumes exceed comfortable on-screen limits.",
    )
    para(
        doc,
        "Progress feedback during long jobs prevents officers from assuming a failure when the "
        "lakehouse is still working. Partial failure in multi-target mode is reported per target "
        "so one bad table does not silently invalidate the entire reconciliation exercise.",
    )
    add_figure(doc, "wf05_analysis_match.png", "Figure 11-5: Analysis match flow")

    heading(doc, "11.6 Access control and data planes", 2)
    para(
        doc,
        "From a governance perspective, SRSE separates who may configure the system from who may "
        "run simulations. Administrators manage registrations and mappings; officers consume "
        "approved assets only.",
    )
    para(
        doc,
        "Operationally, configuration data and analytical data live in different stores. This "
        "separation protects performance and clarifies backup responsibility: configuration can "
        "be exported as JSON; lakehouse data remains owned by the enterprise data platform.",
    )
    add_figure(doc, "wf06_auth_access.png", "Figure 11-6: Roles and dual data planes")

    heading(doc, "12. Outstanding programme activities", 1)
    para(
        doc,
        "The following items require completion or external input before production sign-off. "
        "They do not change the business intent of SRSE but must be tracked at programme level:",
    )
    bullet(doc, "Production SSO claim mapping and UAT with real departmental accounts.")
    bullet(doc, "Confirmation of live golden-layer column names and completion of field mappings.")
    bullet(doc, "Formal SIT/UAT test packs executed against client Presto and DB2 endpoints.")
    bullet(doc, "Finalisation of officer-facing cohort sample workflow where row-level review is required.")

    heading(doc, "13. Data requirements", 1)
    para(doc, "Operational (DB2): field_catalog, field_column_mapping, registered_table, analysis_column_metadata, scenario, scheme linkage.")
    para(doc, "Analytical (Presto/Iceberg): beneficiary and scheme-related tables in Silver and Gold catalogs as registered by admin.")
    para(doc, "Seed data: synthetic beneficiary table for laptop mode (~200k rows) via Docker seed job.")

    heading(doc, "14. Integration interfaces", 1)
    add_table(
        doc,
        ["Interface", "Direction", "Description"],
        [
            ["RajSewadwar SSO", "Inbound", "JWT / session for production auth"],
            ["watsonx.data Presto", "Outbound", "Analytical SQL execution"],
            ["DB2", "Outbound", "JPA operational store"],
            ["SMART portal", "None (phase 1)", "SRSE standalone; future embed possible"],
        ],
    )

    heading(doc, "15. Reference business scenario — Ekal Naari pension", 1)
    para(
        doc,
        "Worked example used for acceptance testing: divorced woman pension with age ≥ 18, "
        "annual income below ₹48,000, domicile in Rajasthan, with exemptions for BPL/Antyodaya and "
        "specified tribal communities (Sahariya, Kathodi, Khairwa). Thresholds align with "
        "consolidated scheme criteria workbook used in SMART stage 1/2.",
    )

    heading(doc, "16. Acceptance criteria (summary)", 1)
    bullet(doc, "Both data planes report UP on /api/health/planes in target environment.")
    bullet(doc, "Officer completes preview for Ekal Naari ruleset and sees count + breakdown.")
    bullet(doc, "Admin registers a Gold table; officer uses its columns in Analysis match.")
    bullet(doc, "Saved scenario compare shows non-zero delta when income threshold changed.")
    bullet(doc, "Analysis CSV download completes for a match exceeding UI row cap.")
    bullet(doc, "Production auth rejects unauthenticated decision calls.")

    heading(doc, "17. Glossary", 1)
    add_table(
        doc,
        ["Term", "Definition"],
        [
            ["Golden Layer", "Curated flat beneficiary dataset in lakehouse Gold catalog"],
            ["Flat catalogue", "Officer-visible fields mapped to single-table columns or expressions"],
            ["Scenario", "Named saved ruleset with stored simulation snapshot"],
            ["Hub (Analysis)", "Central table in multi-target match; one join per target"],
            ["DATA_MODE", "synthetic (local) or live (client lakehouse) configuration switch"],
        ],
    )

    heading(doc, "18. Appendices", 1)
    heading(doc, "Appendix A — Related documents", 2)
    bullet(doc, "docs/SRSE_Technical_Design_Document.docx v1.2 — architecture, service catalog, JSON configuration, Analysis engine")
    bullet(doc, "docs/SRSE_Functional_Workflows.pptx — editable BRD + TDD slides")
    bullet(doc, "docs/CONFIGURATION_GUIDE.md — deployment and environment variables")
    bullet(doc, "CLAUDE.md — locked architecture decisions for implementers")
    bullet(doc, "docs/SRSE_PROJECT_PLAN.xlsx — delivery plan and services inventory")

    heading(doc, "Appendix B — API summary", 2)
    add_table(
        doc,
        ["Area", "Base path", "Key operations"],
        [
            ["Decision", "/api/decision", "preview, cohort, scenarios, compare"],
            ["Analysis", "/api/analysis", "match, match-multi, *.csv, limits"],
            ["Metadata", "/api/metadata", "fields, mappings, analysis columns"],
            ["Lakehouse", "/api/lakehouse", "admin browse, registrations"],
            ["Admin config", "/api/admin/config", "export, import JSON"],
            ["Health", "/api/health", "planes"],
        ],
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
