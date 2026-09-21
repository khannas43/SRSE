"""Build docs/SRSE_Technical_Design_Document.docx (detailed, code-aligned)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "design_assets"
OUT = ROOT / "SRSE_Technical_Design_Document.docx"


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    for r in p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(11)


def mono(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Courier New"
    r.font.size = Pt(9)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
        for p in t.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = val
            for p in t.rows[ri + 1].cells[ci].paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9)
    doc.add_paragraph()


def figure(doc: Document, png: str, caption: str) -> None:
    path = ASSETS / png
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.3))
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in c.runs:
            r.italic = True
            r.font.size = Pt(10)
            r.font.name = "Arial"
    else:
        para(doc, f"[Missing: {png}]")
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Technical Design Document (TDD)")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)
    r.font.name = "Arial"

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run(
        "Scheme Rule Simulation Engine (SRSE)\n"
        "Standalone eligibility simulation and data-reconciliation product"
    )
    sr.font.size = Pt(14)
    sr.font.name = "Arial"

    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = m.add_run(
        f"Version 1.3 | {date.today():%d %B %Y} | Status: For review\n"
        "Editable diagrams: docs/SRSE_Functional_Workflows.pptx\n"
        "Business requirements: docs/SRSE_BRD.docx\n"
        "(First deployed in the Rajasthan SMART programme; readable without prior SMART knowledge.)"
    )
    mr.font.size = Pt(10)
    mr.font.name = "Arial"
    doc.add_page_break()

    # --- 1 ---
    heading(doc, "1. Purpose and audience", 1)
    para(
        doc,
        "This Technical Design Document (TDD) explains what the Scheme Rule Simulation Engine "
        "(SRSE) is, how it is built, and how it behaves at runtime. It is written so that "
        "readers who are not part of the SMART programme—and who may never have seen the "
        "SMART web portal—can still understand SRSE as an independent software product.",
    )
    para(
        doc,
        "SRSE can be deployed against any Presto-compatible lakehouse and operational database "
        "that meets the interfaces described here. Its first production use is in the Government "
        "of Rajasthan SMART initiative, but the architecture deliberately avoids tight coupling "
        "to portal code or release trains.",
    )
    para(
        doc,
        "Primary readers: enterprise architects, security reviewers, integration teams, vendor "
        "engineers, and programme managers. Secondary readers: backend/frontend developers and "
        "lakehouse administrators performing installation and configuration.",
    )

    # --- 2 ---
    heading(doc, "2. About SRSE — system overview", 1)
    para(
        doc,
        "SRSE (Scheme Rule Simulation Engine) is a self-contained web application that helps "
        "government departments estimate how many citizens would qualify for a welfare or "
        "subsidy scheme if eligibility rules or numeric thresholds were changed—for example "
        "lowering an income cap or adding an age band. Instead of requesting custom database "
        "scripts or spreadsheet extracts for every policy question, authorised users compose "
        "rules in a guided user interface and receive counts and breakdowns within minutes.",
    )
    para(
        doc,
        "The product has two major capabilities. Rules simulation lets users build boolean "
        "eligibility conditions from a curated list of citizen attributes (age, income, "
        "domicile, social category, and similar fields), run what-if previews, save named "
        "scenarios, and compare two scenarios side by side. Analysis (record matching) lets "
        "users align two registered datasets—such as a source register and a golden "
        "beneficiary table—using exact or fuzzy column matching, including multi-column name "
        "logic, to support reconciliation and quality checks at scale.",
    )
    para(
        doc,
        "SRSE is delivered as two cooperating containers: a browser-facing application "
        "(Next.js) and an API server (Java Spring Boot). It maintains its own small "
        "operational database for configuration, saved scenarios, and lakehouse registration "
        "metadata. All heavy counting and matching work executes as SQL inside the lakehouse "
        "query engine (Presto over Iceberg tables), so national-scale populations are handled "
        "without loading entire tables into the application server.",
    )
    para(
        doc,
        "Privacy and proportionality are built into the design: routine simulation returns "
        "aggregates only (totals and fixed breakdown dimensions). A separate, strictly capped "
        "API can return a small sample of matching individuals for sanity checks—not the full "
        "eligible population.",
    )
    heading(doc, "2.1 Who uses SRSE", 2)
    table(
        doc,
        ["Role", "Typical tasks"],
        [
            ["Policy / scheme officer", "Adjust thresholds, preview beneficiary counts, save and compare scenarios"],
            ["Data steward / administrator", "Register lakehouse tables, map business fields to physical columns, export/import configuration"],
            ["Integration engineer", "Connect DB2 and Presto endpoints, deploy containers, configure SSO"],
        ],
    )
    heading(doc, "2.2 Deployment modes", 2)
    table(
        doc,
        ["Mode", "Setting", "Behaviour"],
        [
            ["Local synthetic", "DATA_MODE=synthetic", "Docker Presto + seeded Iceberg beneficiary table; field mappings from DB2/seed YAML"],
            ["Client live", "DATA_MODE=live", "On-prem Presto (watsonx.data) + Golden Layer; mappings must use four-part column paths"],
            ["Auth (dev)", "SRSE_AUTH_MODE=mock", "POST /api/auth/mock-login issues JWT for STATE_OFFICER"],
            ["Auth (prod)", "SRSE_AUTH_MODE=rajsewadwar", "RajSewadwar filter (payload parsing seam; fails closed until configured)"],
        ],
    )
    heading(doc, "2.3 Repository layout (implementation)", 2)
    bullet(doc, "backend/ — Java 17 Spring Boot fat JAR (gov.rajasthan.smart.srse.*)")
    bullet(doc, "frontend/ — Next.js App Router pages: /rules, /analysis, /admin456")
    bullet(doc, "docker-compose.yml — local stack (DB2, Presto, MinIO, seed job, both app containers)")
    bullet(doc, "docker-compose.client-dev.yml — profile override for live endpoints")

    # --- 3 ---
    heading(doc, "3. Architecture diagrams", 1)
    para(
        doc,
        "The figures in this section summarise how SRSE fits into a typical deployment. "
        "Each subsection explains the diagram in plain language. Editable versions of these "
        "views (for workshops or customer-specific naming) are in docs/SRSE_Functional_Workflows.pptx, Part B.",
    )

    heading(doc, "3.1 Logical deployment", 2)
    para(
        doc,
        "At the highest level, users interact only with the SRSE web application in a browser. "
        "That application calls the SRSE API over HTTPS. The API is the single integration point "
        "for business logic; it never expects callers to run SQL directly.",
    )
    para(
        doc,
        "The API server connects outward to two data stores. The operational store (DB2 in the "
        "reference deployment) holds SRSE-owned configuration: which fields exist, how they map "
        "to lakehouse columns, which tables officers may use, and saved simulation scenarios. "
        "The analytical connection targets Presto, which reads Iceberg tables that hold "
        "citizen/beneficiary data at scale.",
    )
    para(
        doc,
        "Identity is provided by the organisation's SSO platform (RajSewadwar in the Rajasthan "
        "reference deployment). SRSE validates tokens and roles before allowing rule simulation. "
        "For laptop demos, a built-in mock login can stand in for SSO.",
    )
    figure(doc, "arch01_logical_deployment.png", "Figure 3-1: Containers and external systems")
    para(
        doc,
        "Figure 3-1 emphasises that SRSE is not embedded inside a larger portal container: it "
        "ships and scales as its own frontend and backend pair, which simplifies independent "
        "versioning and security review.",
    )

    heading(doc, "3.2 Dual data plane", 2)
    para(
        doc,
        "A central design choice is the strict separation between operational data and "
        "analytical data. Operational data is small, transactional, and entity-shaped—ideal for "
        "an object-relational mapper (JPA). Analytical workloads are large set-oriented queries "
        "that must stay in the lakehouse engine.",
    )
    para(
        doc,
        "Mixing these planes—for example, attempting to run beneficiary counts through JPA—"
        "would pull too much data into the application tier and break performance guarantees. "
        "SRSE therefore uses Spring Data JPA only for DB2 and raw JDBC (JdbcTemplate) only for Presto.",
    )
    para(
        doc,
        "Officers and administrators never choose which plane to query; the product routes "
        "each request automatically. Saving a scenario touches DB2; previewing a ruleset "
        "touches Presto (after compiling the rule tree to SQL).",
    )
    figure(doc, "arch02_dual_data_plane.png", "Figure 3-2: Operational vs analytical separation")
    para(
        doc,
        "Figure 3-2 should be read as a hard boundary: credentials, backup policies, and "
        "scaling characteristics differ per plane, and only the API layer spans both.",
    )

    heading(doc, "3.3 Backend modules", 2)
    para(
        doc,
        "Inside the API server, responsibilities are grouped into modules (Java packages) "
        "rather than a monolithic service class. REST controllers accept HTTP requests and "
        "perform authentication checks; they delegate to domain services that encode business rules.",
    )
    para(
        doc,
        "The compiler module turns JSON rule trees into SQL. The execution module runs that SQL "
        "on Presto. Metadata and lakehouse modules enforce which columns and tables are legal "
        "to reference. The analysis module builds match queries for the reconciliation UI. "
        "Security and configuration modules wire datasources, timeouts, and authentication filters.",
    )
    para(
        doc,
        "This modular layout keeps the rule engine testable in isolation (unit tests compile "
        "real scheme examples without a browser) and allows future replacement of individual "
        "pieces—such as swapping the REST decision surface for an external rules engine—without "
        "rewriting unrelated code.",
    )
    figure(doc, "arch03_backend_modules.png", "Figure 3-3: Spring Boot package responsibilities")
    para(
        doc,
        "Figure 3-3 is a logical map, not a deployment diagram: all modules run in one JVM "
        "process in the standard fat-JAR deployment.",
    )

    heading(doc, "3.4 Rule preview flow", 2)
    para(
        doc,
        "Rule preview is the most frequent path through SRSE. A user adjusts a threshold in "
        "the Rules screen and asks for an updated count. The browser sends the entire ruleset "
        "JSON to the preview API; nothing is written to the operational database unless the "
        "user explicitly saves a scenario.",
    )
    para(
        doc,
        "The server compiles the ruleset to a parameterised SQL statement. Field names in the "
        "UI (for example annual_income) are resolved to physical lakehouse columns through "
        "mappings stored in DB2. User-entered threshold values are sent to Presto as bound "
        "parameters, not as text embedded in SQL, which prevents injection of officer input.",
    )
    para(
        doc,
        "Presto returns a single aggregate count and, optionally, grouped totals by district, "
        "gender, and age band. Those dimensions are fixed by product policy so that every "
        "scheme uses comparable breakdown views in the UI.",
    )
    figure(doc, "arch04_rule_preview_flow.png", "Figure 3-4: Preview request path")
    para(
        doc,
        "Figure 3-4 highlights that DB2 participates in preview only for metadata lookup, not "
        "for scanning beneficiary rows; all row-level scanning happens in Presto.",
    )

    # --- 4 ---
    heading(doc, "4. Technology stack", 1)
    heading(doc, "4.1 Application tier", 2)
    table(
        doc,
        ["Layer", "Technology", "Version / notes"],
        [
            ["Frontend", "Next.js", "16.1.x"],
            ["UI", "React, HeroUI, Tailwind", "19.2.x / 2.4.x / 4.1.x"],
            ["Client", "TypeScript, Zustand, TanStack Table, Recharts", "5.6 / 5 / 8 / 2"],
            ["Backend", "Spring Boot", "3.3.4"],
            ["Runtime", "Java", "17 (enforced by Maven enforcer)"],
            ["API docs", "springdoc", "/swagger-ui.html"],
        ],
    )
    heading(doc, "4.2 Data tier", 2)
    table(
        doc,
        ["Component", "Technology", "Usage"],
        [
            ["Operational", "DB2 + IBM JCC 11.5.8", "Spring Data JPA repositories"],
            ["Analytical", "PrestoDB 0.297 + presto-jdbc", "JdbcTemplate only — no Hibernate dialect"],
            ["Lakehouse", "Iceberg on object storage", "Silver/Gold beneficiary and related tables"],
        ],
    )

    # --- 5 Services ---
    heading(doc, "5. Service catalog", 1)
    para(
        doc,
        "This catalog lists the HTTP APIs, server-side services, and user-interface modules "
        "that constitute SRSE. It serves as the integration contract for teams connecting "
        "monitoring, API gateways, or alternate frontends. Method and path names match the "
        "implemented Spring Boot controllers unless noted in Section 5.7.",
    )

    heading(doc, "5.1 REST API — decision and scenarios", 2)
    table(
        doc,
        ["Method", "Path", "Controller", "Purpose"],
        [
            ["POST", "/api/decision/preview", "DecisionController", "Compile ruleset + count + optional breakdown (no persist)"],
            ["POST", "/api/decision/cohort", "DecisionController", "Capped row-level cohort sample (only row-level beneficiary API)"],
            ["POST", "/api/decision/scenarios", "DecisionController", "Save scenario + evaluate once + snapshot results"],
            ["GET", "/api/decision/scenarios", "DecisionController", "List scenarios for schemeId"],
            ["GET", "/api/decision/scenarios/{id}", "DecisionController", "Load scenario detail + ruleset JSON"],
            ["GET", "/api/decision/compare", "DecisionController", "Pairwise compare (a, b) — count and breakdown deltas"],
        ],
    )

    heading(doc, "5.2 REST API — analysis and lakehouse (officer)", 2)
    table(
        doc,
        ["Method", "Path", "Controller", "Purpose"],
        [
            ["GET", "/api/analysis/limits", "RecordMatchController", "Server caps (targets, groups, columns)"],
            ["POST", "/api/analysis/match", "RecordMatchController", "Single match — NDJSON stream"],
            ["POST", "/api/analysis/match.csv", "RecordMatchController", "Full CSV re-run from Presto"],
            ["POST", "/api/analysis/match-multi", "RecordMatchController", "Hub vs N targets — NDJSON with per-target progress"],
            ["POST", "/api/analysis/match-multi.csv", "RecordMatchController", "Multi-target CSV (all-or-nothing)"],
            ["GET", "/api/analysis/lakehouse/catalogs/...", "LakehouseCatalogController", "Officer picker: registered cluster subset"],
            ["GET/PUT/DELETE", "/api/analysis/column-metadata", "AnalysisColumnMetadataController", "Per-column business name, fuzzy, compare-as, visibility"],
        ],
    )

    heading(doc, "5.3 REST API — metadata and schemes", 2)
    table(
        doc,
        ["Method", "Path", "Controller", "Purpose"],
        [
            ["GET/POST/PUT/DELETE", "/api/metadata/fields", "FieldCatalogController", "Officer-facing field catalogue CRUD"],
            ["GET/PUT/DELETE", "/api/metadata/mappings", "FieldColumnMappingController", "fieldKey → physicalExpression per DataMode"],
            ["GET/POST", "/api/schemes", "SchemeController", "Welfare scheme tags for scenarios"],
        ],
    )

    heading(doc, "5.4 REST API — administration", 2)
    table(
        doc,
        ["Method", "Path", "Controller", "Purpose"],
        [
            ["GET", "/api/admin/connections", "ConnectionInfoController", "Masked view of JDBC settings"],
            ["PUT", "/api/admin/connections/operational", "ConnectionUpdateController", "Update DB2 pool (+ optional override file)"],
            ["PUT", "/api/admin/connections/analytical", "ConnectionUpdateController", "Update Presto pool"],
            ["GET", "/api/admin/lakehouse/browse/...", "LakehouseAdminController", "Live cluster browse (admin only)"],
            ["GET/POST/PUT/DELETE", "/api/admin/lakehouse/registrations", "LakehouseAdminController", "Officer-visible table registry"],
            ["GET", "/api/admin/config/export", "AdminConfigController", "Download full JSON bundle"],
            ["POST", "/api/admin/config/import", "AdminConfigController", "Upsert bundle; ?testConnections=true default"],
            ["GET", "/api/health/planes", "HealthController", "operational + analytical UP/DOWN"],
            ["POST", "/api/auth/mock-login", "MockJwtIssuer", "Dev-only JWT (when auth-mode=mock)"],
        ],
    )

    heading(doc, "5.5 Application services (backend)", 2)
    table(
        doc,
        ["Service class", "Responsibility"],
        [
            ["RuleCompiler", "Walk AST; emit parameterised SQL + ordered bind values"],
            ["MetadataFieldResolver", "Resolve fieldKey → physical column/expression; Caffeine cache; type coercion"],
            ["ExecutionService", "count(), breakdown(), cohortSample(); applies timeout and cohort cap"],
            ["ScenarioService", "Persist ruleset JSON; record/load/compare snapshots"],
            ["RecordMatchService", "Plan/build match SQL; single-target NDJSON/CSV; display columns; column groups"],
            ["MultiTargetRecordMatchService", "Validate hub; stream meta/progress/row/done; budget enforcement"],
            ["LakehouseBrowseService", "SHOW CATALOGS / information_schema with identifier validation"],
            ["LakehouseRegistryService", "Registered tables; validateColumn (registry + live)"],
            ["FieldColumnMappingService", "Upsert mappings; evict cache on change"],
            ["AdminConfigService", "Export/import JSON bundle; connection merge"],
            ["OperationalConnectionService / AnalyticalConnectionService", "Runtime Hikari reconfiguration"],
            ["ConnectionOverrideStore", "Persist connection-overrides.properties on disk"],
        ],
    )

    heading(doc, "5.6 Frontend modules", 2)
    table(
        doc,
        ["Artifact", "Role"],
        [
            ["app/rules/page.tsx", "Rule builder, preview, scenarios UI"],
            ["app/analysis/page.tsx", "Match criteria, multi-target, streaming grid, CSV"],
            ["app/admin456/page.tsx", "Connections, lakehouse, mappings, JSON backup UI"],
            ["lib/decisionApi.ts", "preview, scenarios, compare; attaches mock JWT when needed"],
            ["lib/analysisApi.ts", "match, match-multi, limits, CSV helpers"],
        ],
    )

    heading(doc, "5.7 Outstanding design and delivery activities", 2)
    para(
        doc,
        "The following items are part of the target product design but require completion, "
        "external input, or formal sign-off before a production cut-over. They are listed "
        "here so programme and integration stakeholders can track dependencies without reading "
        "source code.",
    )
    bullet(
        doc,
        "Production SSO integration — configure RajSewadwar (or equivalent) JWT claim mapping "
        "to the STATE_OFFICER role; security filter shell exists and fails closed until completed.",
    )
    bullet(
        doc,
        "Officer cohort drill-down experience — server API enforces a hard row cap; user-interface "
        "workflow to invoke and interpret cohort samples to be finalised in the Rules module.",
    )
    bullet(
        doc,
        "Live lakehouse field bindings — populate LIVE-mode physical column paths once golden-layer "
        "table and column names are confirmed by the data platform team.",
    )
    bullet(
        doc,
        "Client integration testing and UAT — execute formal test packs against client Presto/DB2 "
        "endpoints and obtain departmental sign-off (see BRD acceptance criteria).",
    )

    # --- 6 Configuration ---
    heading(doc, "6. Runtime configuration (YAML and environment)", 1)
    para(
        doc,
        "Primary file: backend/src/main/resources/application.yml. Spring profiles: local (ddl-auto update for dev), "
        "client-dev (no ddl-auto against real DB2). Environment variables override defaults at container start.",
    )
    heading(doc, "6.1 srse.* application properties", 2)
    table(
        doc,
        ["Property / env", "Default", "Meaning"],
        [
            ["srse.data-mode / DATA_MODE", "synthetic", "Selects FieldResolver bindings (SYNTHETIC vs LIVE)"],
            ["srse.auth-mode / SRSE_AUTH_MODE", "mock", "mock | rajsewadwar"],
            ["srse.datasource.operational.* / SRSE_DB2_*", "local DB2 URL", "JPA datasource (Hikari)"],
            ["srse.datasource.analytical.* / SRSE_PRESTO_*", "local Presto URL", "JdbcTemplate datasource"],
            ["srse.datasource.analytical-ssl.*", "disabled", "SSL trust store — only applied when enabled=true"],
            ["srse.frontend-origins / SRSE_FRONTEND_ORIGINS", "http://localhost:3000", "CORS allow-list"],
            ["srse.guardrails.cohort-cap / SRSE_COHORT_CAP", "1000", "Hard ceiling for cohort drill-down rows"],
            ["srse.guardrails.preview-sample-size / SRSE_PREVIEW_SAMPLE_SIZE", "50", "Default Rules UI sample size (not a second cap; clamped to cohort cap)"],
            ["srse.env-label / SRSE_ENV_LABEL", "(optional)", "Display-only banner override (e.g. UAT)"],
            ["srse.guardrails.query-timeout-seconds", "30", "Presto statement timeout"],
            ["srse.analysis.max-target-sets", "5", "Multi-target limit"],
            ["srse.analysis.multi-match-budget-seconds", "120", "Wall clock budget for multi match stream"],
            ["srse.analysis.max-group-columns", "4", "Columns per side per match group"],
            ["srse.analysis.max-any-of-groups-per-side", "2", "ANY_OF UNNEST cap per side"],
            ["srse.breakdown.age-band-column", "age_band", "Physical column for fixed age_band breakdown"],
            ["srse.connection-override-path", "connection-overrides.properties", "File for admin JDBC overrides"],
        ],
    )
    heading(doc, "6.2 connection-overrides.properties", 2)
    para(
        doc,
        "When an administrator updates JDBC settings via the UI or imports a JSON bundle with "
        "testConnections=false, ConnectionOverrideStore merges keys operational.* and analytical.* "
        "(jdbc-url, username, password, driver-class-name) into this file. "
        "ConnectionOverrideEnvironmentPostProcessor loads the file at boot so overrides survive restarts "
        "without storing credentials inside DB2 (which would be circular for the operational datasource). "
        "Operational plane changes may require application restart (ImportResult.operationalRestartRequired).",
    )

    heading(doc, "6.3 Boot seed (synthetic catalogue)", 2)
    para(
        doc,
        "FieldCatalogSeedRunner upserts field_catalog and field_column_mapping from "
        "backend/src/main/resources/metadata/field-catalog-seed.yml on startup. "
        "LIVE-mode expressions may contain CHANGE_ME until Golden Layer names are supplied.",
    )

    heading(doc, "6.4 Environment display labels vs DATA_MODE", 2)
    para(
        doc,
        "DATA_MODE (synthetic | live) is the only switch that changes which field_column_mapping "
        "rows the resolver uses and which Presto cluster answers analytical queries. It is unchanged "
        "on the wire and in application.yml.",
    )
    bullet(
        doc,
        "The UI banner (Development, UAT, Production (Live), or a custom label via SRSE_ENV_LABEL) "
        "is display-only — it must not be confused with DATA_MODE or with the Admin mapping editor's "
        "SYNTHETIC/LIVE binding-set radios, which select which mapping set is edited, not which "
        "environment the container runs in.",
    )
    bullet(doc, "Local laptops default to DATA_MODE=synthetic; client Dev uses live against on-prem Presto/DB2.")

    heading(doc, "6.5 Lakehouse registration and officer addressing", 2)
    para(
        doc,
        "Lakehouse objects are addressed as catalog.schema.table.column in LIVE bindings and Analysis "
        "SQL. The JDBC URL is catalog-agnostic; every identifier is fully qualified.",
    )
    bullet(
        doc,
        "Admin registration is per table with a mandatory layer tag (BRONZE/SILVER/GOLD/…) for new rows — "
        "layer is a display and registry filter tag, not a hierarchy level in SQL.",
    )
    bullet(
        doc,
        "Officer Analysis cascade: Layer → Catalog → Schema → Table → Column. Layer filters the "
        "registered-table list only; submitted match payloads carry catalog.schema.table, never layer.",
    )
    bullet(
        doc,
        "Legacy registrations with null layer remain reachable under the UNTAGGED sentinel in the "
        "officer cascade (filter only — not stored as a tag name).",
    )
    bullet(doc, "Admin live-browse (LakehouseBrowseService) reaches the whole cluster; officer pickers use LakehouseRegistryService only.")

    # --- 7 JSON bundle ---
    heading(doc, "7. Admin JSON configuration bundle", 1)
    para(
        doc,
        "The admin456 page exports and imports a single JSON document (AdminConfigBundle) via "
        "GET /api/admin/config/export and POST /api/admin/config/import. The bundle uses natural "
        "keys only (no DB2 surrogate ids) so it round-trips across fresh volumes and redeployments.",
    )

    heading(doc, "7.1 Schema version and top-level fields", 2)
    table(
        doc,
        ["JSON field", "Type", "Description"],
        [
            ["schemaVersion", "string", "Must be \"1.0\" (AdminConfigBundle.CURRENT_SCHEMA_VERSION)"],
            ["exportedAt", "ISO-8601 instant", "Export timestamp"],
            ["dataMode", "string", "DATA_MODE at export time (informational)"],
            ["connections", "object", "operational + analytical ConnectionPlane"],
            ["fieldCatalog", "array", "Active catalogue entries"],
            ["fieldColumnMappings", "array", "fieldKey + dataMode + physicalExpression"],
            ["registeredTables", "array", "catalog, schema, table, layer (SILVER/GOLD tag)"],
            ["analysisColumnMetadata", "array", "Four-part column + business metadata"],
            ["schemes", "array", "code, name, description for active schemes"],
        ],
    )

    heading(doc, "7.2 Nested record definitions", 2)
    table(
        doc,
        ["Record", "Fields"],
        [
            ["ConnectionPlane", "jdbcUrl, username, password, driverClassName"],
            ["FieldCatalogEntry", "fieldKey, displayLabel, tier, dataType, groupName, allowedValues[], fuzzyMatchable"],
            ["FieldColumnMappingEntry", "fieldKey, dataMode (SYNTHETIC|LIVE), physicalExpression"],
            ["RegisteredTableEntry", "catalog, schema, table, layer"],
            ["AnalysisColumnMetadataEntry", "catalog, schema, table, column, businessName, fuzzyMatchable, visible, compareAs (AUTO|NUMBER|TEXT)"],
            ["SchemeEntry", "code, name, description"],
        ],
    )

    heading(doc, "7.3 Export behaviour", 2)
    bullet(doc, "Connections: merges live Hikari settings with connection-overrides.properties if present.")
    bullet(doc, "fieldCatalog: active rows only.")
    bullet(doc, "registeredTables: all rows from registered_table entity.")
    bullet(doc, "Response Content-Disposition: srse-admin-config-{date}.json")

    heading(doc, "7.4 Import behaviour and ordering", 2)
    para(doc, "Import runs in a single transaction (mapping cache evicted after). Order of application:")
    bullet(doc, "1. Validate schemaVersion.")
    bullet(doc, "2. connections — if testConnections=true (default query param), test and apply via connection services; else write override file only.")
    bullet(doc, "3. fieldCatalog — upsert by fieldKey.")
    bullet(doc, "4. registeredTables — LakehouseRegistryService.importRegistration per entry.")
    bullet(doc, "5. fieldColumnMappings — skip blank physicalExpression; require existing active field.")
    bullet(doc, "6. analysisColumnMetadata — upsert by qualified column.")
    bullet(doc, "7. schemes — upsert by code.")
    para(doc, "Response body: ImportResult with counts and lists of imported keys; operationalRestartRequired when operational JDBC changed.")

    heading(doc, "7.5 Security and operational notes", 2)
    bullet(doc, "JSON contains database passwords in plain text — treat exports as confidential; restrict file storage.")
    bullet(doc, "Import does not delete catalogue entries absent from the file; it upserts listed keys only.")
    bullet(doc, "Re-import on production should use testConnections=true in a maintenance window.")

    heading(doc, "7.6 Example skeleton (illustrative)", 2)
    mono(
        doc,
        '{\n'
        '  "schemaVersion": "1.0",\n'
        '  "exportedAt": "2026-09-20T10:00:00Z",\n'
        '  "dataMode": "live",\n'
        '  "connections": {\n'
        '    "operational": { "jdbcUrl": "jdbc:db2://...", "username": "...", "password": "***", "driverClassName": "com.ibm.db2.jcc.DB2Driver" },\n'
        '    "analytical": { "jdbcUrl": "jdbc:presto://...", "username": "srse", "password": "", "driverClassName": "com.facebook.presto.jdbc.PrestoDriver" }\n'
        '  },\n'
        '  "fieldCatalog": [ { "fieldKey": "age_years", "displayLabel": "Age (years)", "tier": "TIER1", "dataType": "NUMBER", ... } ],\n'
        '  "fieldColumnMappings": [ { "fieldKey": "age_years", "dataMode": "LIVE", "physicalExpression": "iceberg.gold.beneficiary.age_years" } ],\n'
        '  "registeredTables": [ { "catalog": "iceberg", "schema": "gold", "table": "beneficiary", "layer": "GOLD" } ],\n'
        '  "analysisColumnMetadata": [ { "catalog": "...", "column": "full_name", "compareAs": "AUTO", "visible": true, ... } ],\n'
        '  "schemes": [ { "code": "EKAL_NAARI", "name": "Ekal Naari Pension", "description": "..." } ]\n'
        '}',
    )

    # --- 8 Compiler ---
    heading(doc, "8. Rule compiler and execution pipeline", 1)
    heading(doc, "8.1 AST model (compiler/Ast.java)", 2)
    para(
        doc,
        "Rulesets serialise as JSON polymorphic nodes: GROUP (AND/OR + children) and PREDICATE "
        "(fieldKey, operator, value). Operators include EQ, NE, comparisons, IN, NOT_IN, BETWEEN, "
        "IS_TRUE/FALSE, IS_NULL/NOT_NULL, FUZZY_MATCH. No join nodes — flat catalogue only.",
    )
    heading(doc, "8.2 Compilation and execution", 2)
    bullet(doc, "RuleCompiler resolves each fieldKey through MetadataFieldResolver for current DATA_MODE.")
    bullet(doc, "Emitted SQL uses ? placeholders; officer literals never appear in SQL text.")
    bullet(doc, "ExecutionService runs COUNT query and optional GROUP BY breakdown (district, gender, age_band).")
    bullet(doc, "cohortSample applies effectiveCohortLimit — caller cannot exceed SRSE_COHORT_CAP.")
    bullet(doc, "Query timeout from GuardrailProperties applied to JdbcTemplate execution.")

    heading(doc, "8.3 Preview count, cohort sample, and row shape", 2)
    para(
        doc,
        "Preview (/api/decision/preview) returns aggregates only — total count and optional fixed "
        "breakdown dimensions (district, gender, age_band). Cohort (/api/decision/cohort) is the "
        "only endpoint that returns row-level beneficiary data, hard-capped by SRSE_COHORT_CAP; "
        "the response states the limit applied and whether the sample was truncated.",
    )
    bullet(
        doc,
        "The Rules UI default sample size is SRSE_PREVIEW_SAMPLE_SIZE (default 50), clamped to "
        "SRSE_COHORT_CAP — it is not a separate server ceiling.",
    )
    bullet(
        doc,
        "Cohort and preview sample queries deliberately use SELECT * on the beneficiary row for "
        "departmental users who need to sanity-check real records; narrowing projection requires an "
        "explicit product decision.",
    )

    heading(doc, "8.4 Scheme official criteria vs officer scenarios", 2)
    para(
        doc,
        "Each scheme may nominate one saved scenario as its official template (scheme.template_scenario_id). "
        "Officers loading that template into the Rules builder start from official criteria but Save always "
        "creates a new scenario — existing scenario rulesets are write-once; there is no PUT/PATCH on "
        "/api/decision/scenarios/**. Only SRSE_ADMIN may change which scenario is the template "
        "(PUT /api/schemes/{id}/template).",
    )

    # --- 9 Analysis ---
    heading(doc, "9. Analysis design — record matching engine", 1)
    para(
        doc,
        "The Analysis capability answers a different question from Rules simulation. Where "
        "Rules asks how many people satisfy a policy ruleset, Analysis asks which rows in "
        "dataset A correspond to rows in dataset B—and where they differ. Typical uses include "
        "reconciling a departmental register against a golden beneficiary table, finding "
        "duplicates under spelling variation, or validating that two lakehouse layers (often "
        "labelled Silver and Gold) refer to the same individuals.",
    )
    para(
        doc,
        "From a stakeholder perspective, Analysis is a guided join builder: the user selects "
        "two registered tables, defines one or more match criteria (exact equality or fuzzy "
        "name match), optionally adds columns that should appear in the result but not affect "
        "matching, and runs the job. Results stream back progressively for responsiveness; "
        "very large result sets are downloaded as CSV generated directly from Presto so the "
        "user receives the complete match set, not merely what the browser could hold in memory.",
    )
    para(
        doc,
        "The Analysis engine is implemented as server-side services (RecordMatchService for "
        "single source–target pairs, MultiTargetRecordMatchService when one hub table is "
        "compared against several targets in one session). Both share the same safety model: "
        "only administrator-registered tables and live-confirmed columns may appear in SQL; "
        "cross-type comparisons follow documented coercion rules so that, for example, numeric "
        "IDs stored as text still match sensibly.",
    )
    heading(doc, "9.1 What users can do in Analysis", 2)
    bullet(doc, "Match on one column per side, or on column groups (combined name fields, or match-if-any-of several columns).")
    bullet(doc, "Choose exact match or fuzzy match (string similarity score exposed in results).")
    bullet(doc, "Add display-only columns for context (e.g. show district on both sides without using it in the join key).")
    bullet(doc, "Run multi-target mode: one central (hub) table matched independently to many target tables, with combined progress reporting.")
    bullet(doc, "Export full results to CSV when on-screen grids would be impractical at crore scale.")

    heading(doc, "9.2 Engine pipeline (technical)", 2)
    para(
        doc,
        "Each run begins with validation: tables must be registered, columns must still exist in "
        "Presto, and group sizes must respect configured caps (maximum groups, ANY_OF groups, "
        "multi-target count, and wall-clock budget for long multi-target streams).",
    )
    para(
        doc,
        "The planner constructs a two-table SQL query (never an N-way join for multi-target—"
        "each target is a separate query plan). Column groups emit specialised SQL: COMBINE "
        "folds multiple columns into one comparable string; ANY_OF uses an UNNEST pivot so "
        "Presto retains efficient join plans. Identifiers and catalog names pass allow-list checks "
        "before being embedded in SQL text; literal match values use parameter binding.",
    )
    para(
        doc,
        "Execution streams newline-delimited JSON events to the client: metadata first, then "
        "progress, row payloads, completion or per-target errors. CSV endpoints discard the "
        "streaming UI path and pipe Presto output straight to the download. Emitted SQL shapes "
        "are validated in automated tests against the Presto 0.297 grammar parser to reduce "
        "risk of runtime syntax surprises.",
    )

    heading(doc, "9.3 Presentation-layer behaviour", 2)
    bullet(doc, "Browser grid and charts render up to 10,000 rows; beyond that the UI directs users to CSV.")
    bullet(doc, "NDJSON consumption stops after 200,000 rows for browser stability; displayed counts may show as a lower bound while CSV remains complete.")
    bullet(doc, "Multi-target CSV export is all-or-nothing—one failed target aborts the file so partial exports are not mistaken for full reconciliation.")

    heading(doc, "9.4 Two-table join types (single match)", 2)
    para(
        doc,
        "POST /api/analysis/match accepts joinType INNER (default), LEFT, RIGHT, or FULL for exactly "
        "two registered tables. Self-join is supported by picking the same table on both sides — no "
        "separate mode. POST /api/analysis/match.sql plans and returns display SQL without executing. "
        "Outer-join predicate routing (fuzzy and ANY_OF in ON, not WHERE) follows RecordMatchService; "
        "dedup is rejected with RIGHT/FULL.",
    )
    bullet(doc, "Multi-target mode (match-multi) remains INNER only — N independent hub↔target joins; arbitrary N-way joins are deferred.")

    heading(doc, "9.5 Multi-target join canvas (Phase A)", 2)
    para(
        doc,
        "The Analysis UI offers a join canvas view over the existing MultiTargetRecordMatchRequest — "
        "no new SQL or backend records. Hub node centre, target nodes, ordered edges (one per join "
        "criterion). Serialisation derives hubCriteria and each target's joinCriteria from the same "
        "ordered slot list so index alignment cannot drift. Per-target SQL appears on progress events "
        "with phase started (never on meta); planning failures may emit error without started; shared "
        "multiMatchBudgetSeconds may skip later targets with reason time budget.",
    )

    heading(doc, "9.6 Join-key suggestions", 2)
    para(
        doc,
        "POST /api/analysis/suggest-keys returns metadata-first column pair hints (name/type heuristics). "
        "With probe=true the service samples the source table (TABLESAMPLE BERNOULLI), scans the target "
        "in full per pair, and ranks overlap — bounded by SRSE_ANALYSIS_MAX_PROBED_PAIRS and query timeout.",
    )

    # --- 10 Data model ---
    heading(doc, "10. Operational data model (DB2 / JPA)", 1)
    table(
        doc,
        ["Entity", "Purpose"],
        [
            ["FieldCatalogEntry", "Officer-visible field definitions (tier, data type, allowed values)"],
            ["FieldColumnMapping", "Composite key: fieldKey + dataMode → physicalExpression"],
            ["RegisteredTable", "Admin-registered lakehouse table (catalog, schema, table, layer tag)"],
            ["AnalysisColumnMetadata", "Overrides per lakehouse column for Analysis UI"],
            ["Scenario", "Named ruleset JSON + linked scheme ids + result snapshot fields"],
            ["Scheme", "Scheme code/name; optional templateScenarioId for official criteria"],
        ],
    )
    para(doc, "Scenario ruleset stored as JSON using Ast.Node polymorphism for round-trip with the UI.")

    # --- 11 Security ---
    heading(doc, "11. Security and authentication", 1)
    para(
        doc,
        "Production access uses RajSewadwar SSO mapped to Spring authorities. SRSE_ADMIN gates "
        "live lakehouse browse, registration, field catalogue, JSON backup, and scheme template "
        "assignment. STATE_OFFICER reaches decision preview/scenarios, officer Analysis APIs, and "
        "read-only metadata needed to build rules.",
    )
    bullet(doc, "SecurityConfig enforces per-path roles; SecurityConfigRbacTest guards drift.")
    bullet(
        doc,
        "SRSE_AUTH_MODE=mock is a development seam only (?role=admin mints tokens) — not production "
        "access control. Production role mapping binds in RajSewadwarAuthenticationFilter.grantedAuthoritiesFromSsoRoles "
        "(Aadhaar OTP / project dev team integration point).",
    )
    bullet(doc, "CORS: srse.frontend-origins must match browser origin (not Docker internal hostname).")
    bullet(doc, "LakehouseIdentifiers + SHOW CATALOGS allow-list before interpolating catalog/schema in SQL.")
    bullet(doc, "Analysis: registered table + live column existence required before column reaches SQL.")

    # --- 12 Locked ---
    heading(doc, "12. Locked architectural decisions", 1)
    table(
        doc,
        ["Decision", "Rationale"],
        [
            ["Standalone microservice", "Decoupled portal release cycle"],
            ["Two data planes", "JPA for metadata; JdbcTemplate for Presto set queries"],
            ["PrestoDB 0.297 driver", "watsonx.data lineage — not Trino"],
            ["Flat catalogue / no JOINs in compiler", "Tier-3 fields pre-materialised upstream"],
            ["REST decision seam", "DMN-shaped API; swappable authoring backend"],
            ["Push-down execution", "Counts never pull full populations into JVM"],
        ],
    )

    # --- 13 Related ---
    heading(doc, "13. Related documents", 1)
    bullet(doc, "docs/SRSE_BRD.docx")
    bullet(doc, "docs/SRSE_Functional_Workflows.pptx")
    bullet(doc, "docs/CONFIGURATION_GUIDE.md")
    bullet(doc, "CLAUDE.md")
    bullet(doc, "docs/SRSE_PROJECT_PLAN.xlsx")

    # --- 14 History ---
    heading(doc, "14. Document history", 1)
    table(
        doc,
        ["Version", "Date", "Change"],
        [
            ["1.0", "2026-09-20", "Initial TDD with architecture figures and stack tables"],
            ["1.1", "2026-09-20", "Detailed services, YAML/env config, JSON bundle spec; removed guardrails section"],
            [
                "1.2",
                "2026-09-20",
                "Stakeholder-oriented §2; narrative for §3 diagrams; Service catalog wording; Analysis engine intro; §5.7 delivery activities",
            ],
            [
                "1.3",
                date.today().isoformat(),
                "Lakehouse layer filter; env labels; preview/cohort sample; join types; join canvas Phase A; suggest-keys probe; RBAC; scheme templates",
            ],
        ],
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
